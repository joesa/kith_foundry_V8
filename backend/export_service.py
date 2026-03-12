"""
Export Service — renders C-Suite analyses and Artifacts to .md, .docx, or .pdf bytes.

Usage (from export_api.py):
    from export_service import render_csuite, render_artifact, render_all_artifacts
    content, media_type, filename = render_csuite(project, analyses, fmt)
"""
from __future__ import annotations

import io
import textwrap
from typing import Literal

Format = Literal["md", "docx", "pdf"]

# ── Sanitize for PDF (Latin-1 core fonts) ─────────────────────────────────────
_REPLACE = {
    "\u2014": "--", "\u2013": "-", "\u2018": "'", "\u2019": "'",
    "\u201c": '"',  "\u201d": '"',  "\u2022": "*", "\u00d7": "x",
    "\u2265": ">=", "\u2264": "<=", "\u2248": "~=", "\u2192": "->",
    "\u00b1": "+/-",
}

def _safe(text: str) -> str:
    for ch, rep in _REPLACE.items():
        text = text.replace(ch, rep)
    return text.encode("latin-1", errors="replace").decode("latin-1")


def _s(v) -> str:
    """Coerce any value to a non-None string."""
    return str(v) if v is not None else ""


# ─────────────────────────────────────────────────────────────────────────────
# MARKDOWN helpers
# ─────────────────────────────────────────────────────────────────────────────

def _list_md(items) -> str:
    if not items:
        return "_None provided._\n"
    return "".join(f"- {item}\n" for item in items)


def _csuite_analysis_md(project_name: str, analyses: list) -> str:
    ROLE_LABELS = {
        "ceo": "CEO — Strategic Vision",
        "cto": "CTO — Technical Feasibility",
        "cfo": "CFO — Financial Analysis",
        "cmo": "CMO — Go-to-Market",
        "cpo": "CPO — Product Strategy",
        "coo": "COO — Operations",
        "cdo": "CDO — Design & Brand",
    }
    lines = [
        f"# C-Suite Analysis Report\n",
        f"**Project:** {project_name}\n",
        "---\n",
    ]
    # Overall summary
    complete = [a for a in analyses if a.get("status") == "complete" and a.get("score") is not None]
    if complete:
        avg = round(sum(a["score"] for a in complete) / len(complete))
        lines.append(f"## Overall Score: {avg}/100\n")

    for a in analyses:
        role = a.get("agent_role", a.get("role", ""))
        label = ROLE_LABELS.get(role, role.upper())
        data = a.get("analysis") or {}
        status = a.get("status", "pending")
        lines.append(f"## {label}\n")
        if status != "complete" or not data:
            lines.append(f"_Status: {status}_\n\n")
            continue
        score = data.get("score")
        verdict = data.get("verdict", "")
        lines.append(f"**Score:** {score}/100  |  **Verdict:** {verdict.upper()}\n\n")
        lines.append(f"### Recommendation\n{_s(data.get('recommendation'))}\n\n")
        if data.get("deep_analysis"):
            lines.append(f"### Deep Analysis\n{_s(data.get('deep_analysis'))}\n\n")
        if data.get("strengths"):
            lines.append(f"### Strengths\n{_list_md(data['strengths'])}\n")
        if data.get("risks"):
            lines.append(f"### Risks\n{_list_md(data['risks'])}\n")
        if data.get("priority_actions"):
            lines.append(f"### Priority Actions\n{_list_md(data['priority_actions'])}\n")
        if data.get("suggestions"):
            lines.append(f"### Suggestions\n{_list_md(data['suggestions'])}\n")
        if data.get("key_metrics"):
            lines.append(f"### Key Metrics\n{_list_md(data['key_metrics'])}\n")
        if data.get("timeline"):
            lines.append(f"### Timeline\n{_s(data.get('timeline'))}\n\n")
        if data.get("competitive_note"):
            lines.append(f"### Competitive Note\n{_s(data.get('competitive_note'))}\n\n")
        lines.append("---\n")
    return "\n".join(lines)


def _artifact_md(project_name: str, title: str, content) -> str:
    lines = [f"# {title}\n", f"**Project:** {project_name}\n", "---\n"]
    if isinstance(content, dict):
        for key, value in content.items():
            heading = key.replace("_", " ").title()
            lines.append(f"## {heading}\n")
            if isinstance(value, list):
                lines.append(_list_md(value))
            elif isinstance(value, dict):
                for k2, v2 in value.items():
                    lines.append(f"### {k2.replace('_', ' ').title()}\n")
                    if isinstance(v2, list):
                        lines.append(_list_md(v2))
                    else:
                        lines.append(f"{_s(v2)}\n\n")
            else:
                lines.append(f"{_s(value)}\n\n")
    elif isinstance(content, str):
        lines.append(content)
    return "\n".join(lines)


def _all_artifacts_md(project_name: str, artifacts: list, separator: bool = True) -> str:
    parts = [f"# Project Artifacts\n\n**Project:** {project_name}\n\n---\n\n"]
    for art in artifacts:
        if art.get("status") != "complete":
            continue
        parts.append(_artifact_md(project_name, art.get("title", "Artifact"), art.get("content")))
        if separator:
            parts.append("\n\n---\n\n")
    return "".join(parts)


# ─────────────────────────────────────────────────────────────────────────────
# DOCX helpers
# ─────────────────────────────────────────────────────────────────────────────

def _docx_shade_row(row, hex_color: str = "D9E1F2"):
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    for cell in row.cells:
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), hex_color)
        tcPr.append(shd)


def _docx_new() -> "Document":
    from docx import Document
    from docx.shared import Inches
    doc = Document()
    for section in doc.sections:
        section.top_margin    = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin   = Inches(1.25)
        section.right_margin  = Inches(1.25)
    return doc


def _docx_title_block(doc, title: str, subtitle: str):
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    t = doc.add_heading(title, level=0)
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph(subtitle)
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.runs[0].font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    doc.add_paragraph()


def _docx_bullets(doc, items):
    for item in (items or []):
        doc.add_paragraph(str(item), style="List Bullet")


def _docx_kv_table(doc, rows: list[tuple[str, str]]):
    from docx.shared import Pt, RGBColor
    tbl = doc.add_table(rows=0, cols=2)
    tbl.style = "Table Grid"
    for i, (k, v) in enumerate(rows):
        row = tbl.add_row()
        if i % 2 == 0:
            _docx_shade_row(row, "EDF2FF")
        row.cells[0].text = k
        row.cells[0].paragraphs[0].runs[0].font.bold = True
        row.cells[0].paragraphs[0].runs[0].font.size = Pt(9)
        row.cells[1].text = v
        row.cells[1].paragraphs[0].runs[0].font.size = Pt(9)
    doc.add_paragraph()


def _csuite_analysis_docx(project_name: str, analyses: list) -> bytes:
    from docx.shared import Pt, RGBColor
    ROLE_LABELS = {
        "ceo": "CEO — Strategic Vision",    "cto": "CTO — Technical Feasibility",
        "cfo": "CFO — Financial Analysis",  "cmo": "CMO — Go-to-Market",
        "cpo": "CPO — Product Strategy",    "coo": "COO — Operations",
        "cdo": "CDO — Design & Brand",
    }
    doc = _docx_new()
    _docx_title_block(doc, "C-Suite Analysis Report", f"Project: {project_name}")

    complete = [a for a in analyses if a.get("status") == "complete" and a.get("score") is not None]
    if complete:
        avg = round(sum(a["score"] for a in complete) / len(complete))
        doc.add_heading(f"Overall Score: {avg}/100", level=1)

    for a in analyses:
        role = a.get("agent_role", a.get("role", ""))
        label = ROLE_LABELS.get(role, role.upper())
        data = a.get("analysis") or {}
        status = a.get("status", "pending")
        doc.add_heading(label, level=1)
        if status != "complete" or not data:
            doc.add_paragraph(f"Status: {status}")
            continue
        score = data.get("score")
        verdict = data.get("verdict", "")
        doc.add_heading("Recommendation", level=2)
        doc.add_paragraph(f"Score: {score}/100  |  Verdict: {verdict.upper()}")
        doc.add_paragraph(_s(data.get("recommendation", "")))
        if data.get("deep_analysis"):
            doc.add_heading("Deep Analysis", level=2)
            doc.add_paragraph(_s(data["deep_analysis"]))
        for section, items in [
            ("Strengths", data.get("strengths")),
            ("Risks", data.get("risks")),
            ("Priority Actions", data.get("priority_actions")),
            ("Suggestions", data.get("suggestions")),
            ("Key Metrics", data.get("key_metrics")),
        ]:
            if items:
                doc.add_heading(section, level=2)
                _docx_bullets(doc, items)
        if data.get("timeline"):
            doc.add_heading("Timeline", level=2)
            doc.add_paragraph(_s(data["timeline"]))
        if data.get("competitive_note"):
            doc.add_heading("Competitive Note", level=2)
            doc.add_paragraph(_s(data["competitive_note"]))

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _artifact_docx(project_name: str, title: str, content) -> bytes:
    doc = _docx_new()
    _docx_title_block(doc, title, f"Project: {project_name}")
    _artifact_content_to_docx(doc, content)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _artifact_content_to_docx(doc, content):
    if isinstance(content, dict):
        for key, value in content.items():
            heading = key.replace("_", " ").title()
            doc.add_heading(heading, level=2)
            if isinstance(value, list):
                _docx_bullets(doc, value)
            elif isinstance(value, dict):
                rows = [(k2.replace("_", " ").title(), _s(v2) if not isinstance(v2, list) else ", ".join(str(i) for i in v2))
                        for k2, v2 in value.items()]
                _docx_kv_table(doc, rows)
            else:
                doc.add_paragraph(_s(value))
    elif isinstance(content, str):
        doc.add_paragraph(content)


def _all_artifacts_docx(project_name: str, artifacts: list) -> bytes:
    doc = _docx_new()
    _docx_title_block(doc, "Project Artifacts", f"Project: {project_name}")
    for art in artifacts:
        if art.get("status") != "complete":
            continue
        doc.add_heading(art.get("title", "Artifact"), level=1)
        _artifact_content_to_docx(doc, art.get("content"))
        doc.add_page_break()
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ─────────────────────────────────────────────────────────────────────────────
# PDF helpers
# ─────────────────────────────────────────────────────────────────────────────

def _make_pdf(title: str, subtitle: str):
    from fpdf import FPDF
    from fpdf.enums import XPos, YPos

    _title_ref = [title]

    class PDF(FPDF):
        def header(self):
            self.set_font("Helvetica", "B", 9)
            self.set_text_color(120, 120, 120)
            self.cell(0, 8, _safe(_title_ref[0])[:80], align="L",
                      new_x=XPos.LMARGIN, new_y=YPos.NEXT)
            self.ln(1)

        def footer(self):
            self.set_y(-15)
            self.set_font("Helvetica", "", 8)
            self.set_text_color(150, 150, 150)
            self.cell(0, 8, f"Page {self.page_no()}", align="C")

    pdf = PDF()
    pdf.set_auto_page_break(auto=True, margin=20)
    pdf.set_margins(20, 20, 20)
    pdf.add_page()

    from fpdf.enums import XPos, YPos

    # Title block
    pdf.set_font("Helvetica", "B", 16)
    pdf.set_text_color(30, 30, 30)
    pdf.multi_cell(0, 9, _safe(title), align="C", new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.set_font("Helvetica", "", 10)
    pdf.set_text_color(100, 100, 100)
    pdf.multi_cell(0, 7, _safe(subtitle), align="C", new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.ln(5)

    return pdf


def _pdf_h1(pdf, text: str):
    from fpdf.enums import XPos, YPos
    pdf.set_font("Helvetica", "B", 13)
    pdf.set_text_color(30, 60, 120)
    pdf.multi_cell(0, 8, _safe(text), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.ln(1)


def _pdf_h2(pdf, text: str):
    from fpdf.enums import XPos, YPos
    pdf.set_font("Helvetica", "B", 10)
    pdf.set_text_color(50, 50, 50)
    pdf.multi_cell(0, 7, _safe(text), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.ln(1)


def _pdf_para(pdf, text: str):
    from fpdf.enums import XPos, YPos
    pdf.set_font("Helvetica", "", 9)
    pdf.set_text_color(30, 30, 30)
    if text:
        pdf.multi_cell(0, 6, _safe(str(text)), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.ln(2)


def _pdf_bullets(pdf, items):
    from fpdf.enums import XPos, YPos
    pdf.set_font("Helvetica", "", 9)
    pdf.set_text_color(30, 30, 30)
    for item in (items or []):
        pdf.set_x(pdf.get_x() + 5)
        pdf.multi_cell(0, 6, _safe(f"* {item}"), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.ln(2)


def _pdf_kv_rows(pdf, pairs: list[tuple[str, str]]):
    from fpdf.enums import XPos, YPos
    usable = pdf.w - pdf.l_margin - pdf.r_margin
    col1 = usable * 0.35
    col2 = usable * 0.65
    pdf.set_fill_color(46, 64, 87)
    pdf.set_text_color(255, 255, 255)
    pdf.set_font("Helvetica", "B", 8)
    pdf.cell(col1, 7, "Field", border=1, fill=True)
    pdf.cell(col2, 7, "Value", border=1, fill=True)
    pdf.ln()
    pdf.set_font("Helvetica", "", 8)
    for i, (k, v) in enumerate(pairs):
        if i % 2 == 0:
            pdf.set_fill_color(237, 242, 255)
        else:
            pdf.set_fill_color(255, 255, 255)
        pdf.set_text_color(30, 30, 30)
        pdf.cell(col1, 6, _safe(str(k))[:40], border=1, fill=True)
        pdf.cell(col2, 6, _safe(str(v))[:60], border=1, fill=True)
        pdf.ln()
    pdf.ln(3)


def _artifact_content_to_pdf(pdf, content):
    if isinstance(content, dict):
        for key, value in content.items():
            _pdf_h2(pdf, key.replace("_", " ").title())
            if isinstance(value, list):
                _pdf_bullets(pdf, value)
            elif isinstance(value, dict):
                pairs = [(k2.replace("_", " ").title(),
                          ", ".join(str(i) for i in v2) if isinstance(v2, list) else _s(v2))
                         for k2, v2 in value.items()]
                _pdf_kv_rows(pdf, pairs)
            else:
                _pdf_para(pdf, _s(value))
    elif isinstance(content, str):
        _pdf_para(pdf, content)


def _csuite_analysis_pdf(project_name: str, analyses: list) -> bytes:
    ROLE_LABELS = {
        "ceo": "CEO -- Strategic Vision",    "cto": "CTO -- Technical Feasibility",
        "cfo": "CFO -- Financial Analysis",  "cmo": "CMO -- Go-to-Market",
        "cpo": "CPO -- Product Strategy",    "coo": "COO -- Operations",
        "cdo": "CDO -- Design & Brand",
    }
    pdf = _make_pdf("C-Suite Analysis Report", f"Project: {project_name}")

    complete = [a for a in analyses if a.get("status") == "complete" and a.get("score") is not None]
    if complete:
        avg = round(sum(a["score"] for a in complete) / len(complete))
        _pdf_h1(pdf, f"Overall Score: {avg}/100")

    for a in analyses:
        role = a.get("agent_role", a.get("role", ""))
        label = ROLE_LABELS.get(role, role.upper())
        data = a.get("analysis") or {}
        status = a.get("status", "pending")
        _pdf_h1(pdf, label)
        if status != "complete" or not data:
            _pdf_para(pdf, f"Status: {status}")
            continue
        score = data.get("score")
        verdict = _s(data.get("verdict", ""))
        _pdf_para(pdf, f"Score: {score}/100   |   Verdict: {verdict.upper()}")
        _pdf_h2(pdf, "Recommendation")
        _pdf_para(pdf, _s(data.get("recommendation", "")))
        if data.get("deep_analysis"):
            _pdf_h2(pdf, "Deep Analysis")
            _pdf_para(pdf, _s(data["deep_analysis"]))
        for section, items in [
            ("Strengths", data.get("strengths")),
            ("Risks", data.get("risks")),
            ("Priority Actions", data.get("priority_actions")),
            ("Suggestions", data.get("suggestions")),
            ("Key Metrics", data.get("key_metrics")),
        ]:
            if items:
                _pdf_h2(pdf, section)
                _pdf_bullets(pdf, items)
        if data.get("timeline"):
            _pdf_h2(pdf, "Timeline")
            _pdf_para(pdf, _s(data["timeline"]))
        if data.get("competitive_note"):
            _pdf_h2(pdf, "Competitive Note")
            _pdf_para(pdf, _s(data["competitive_note"]))

    buf = io.BytesIO()
    pdf.output(buf)
    return buf.getvalue()


def _artifact_pdf(project_name: str, title: str, content) -> bytes:
    pdf = _make_pdf(title, f"Project: {project_name}")
    _artifact_content_to_pdf(pdf, content)
    buf = io.BytesIO()
    pdf.output(buf)
    return buf.getvalue()


def _all_artifacts_pdf(project_name: str, artifacts: list) -> bytes:
    pdf = _make_pdf("Project Artifacts", f"Project: {project_name}")
    for art in artifacts:
        if art.get("status") != "complete":
            continue
        _pdf_h1(pdf, art.get("title", "Artifact"))
        _artifact_content_to_pdf(pdf, art.get("content"))
        pdf.add_page()
    buf = io.BytesIO()
    pdf.output(buf)
    return buf.getvalue()


# ─────────────────────────────────────────────────────────────────────────────
# Public API
# ─────────────────────────────────────────────────────────────────────────────

def render_csuite(project_name: str, analyses: list, fmt: Format) -> tuple[bytes, str, str]:
    """Returns (content_bytes, media_type, filename)."""
    slug = project_name.lower().replace(" ", "_")[:40]
    if fmt == "md":
        md = _csuite_analysis_md(project_name, analyses)
        return md.encode(), "text/markdown", f"{slug}_csuite.md"
    if fmt == "docx":
        return _csuite_analysis_docx(project_name, analyses), \
               "application/vnd.openxmlformats-officedocument.wordprocessingml.document", \
               f"{slug}_csuite.docx"
    # pdf
    return _csuite_analysis_pdf(project_name, analyses), "application/pdf", f"{slug}_csuite.pdf"


def render_artifact(project_name: str, title: str, content, fmt: Format) -> tuple[bytes, str, str]:
    """Returns (content_bytes, media_type, filename)."""
    slug = title.lower().replace(" ", "_")[:40]
    pslug = project_name.lower().replace(" ", "_")[:30]
    if fmt == "md":
        md = _artifact_md(project_name, title, content)
        return md.encode(), "text/markdown", f"{pslug}_{slug}.md"
    if fmt == "docx":
        return _artifact_docx(project_name, title, content), \
               "application/vnd.openxmlformats-officedocument.wordprocessingml.document", \
               f"{pslug}_{slug}.docx"
    return _artifact_pdf(project_name, title, content), "application/pdf", f"{pslug}_{slug}.pdf"


def render_all_artifacts(project_name: str, artifacts: list, fmt: Format) -> tuple[bytes, str, str]:
    """Returns (content_bytes, media_type, filename)."""
    pslug = project_name.lower().replace(" ", "_")[:40]
    if fmt == "md":
        md = _all_artifacts_md(project_name, artifacts)
        return md.encode(), "text/markdown", f"{pslug}_artifacts.md"
    if fmt == "docx":
        return _all_artifacts_docx(project_name, artifacts), \
               "application/vnd.openxmlformats-officedocument.wordprocessingml.document", \
               f"{pslug}_artifacts.docx"
    return _all_artifacts_pdf(project_name, artifacts), "application/pdf", f"{pslug}_artifacts.pdf"
